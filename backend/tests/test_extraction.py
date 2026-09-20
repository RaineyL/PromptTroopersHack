import json
import os
from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest
from unittest.mock import Mock, patch

from docx import Document
from fastapi import HTTPException
from fastapi.testclient import TestClient
from openpyxl import Workbook
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from app.main import app
from app.schemas.classification import Email
from app.services.extraction.catalog import ExtractionCatalog
from app.services.extraction.documents import DocumentReadError, read_document
from app.services.extraction.fields import document_kind, extract_fields
from app.services.extraction.pipeline import extract_email
from app.services.extraction.pdf import ocr_page
from app.services.inbox import InboxClient

BL = b'''BILL OF LADING (DRAFT)\nSHIPPER: Example Exporter\nCONSIGNEE: Example Receiver\nNotify: Example Agent\nPort of Loading (POL): PORT KLANG\nPOD: CALLAO\nContainer Count: 2 x 40'HC\nGross Wt (kgs): 21,577 KG\n'''
SI = BL.replace(b'BILL OF LADING (DRAFT)', b'SHIPPING INSTRUCTION')


class ExtractionTests(unittest.TestCase):
    def test_catalog_only_uses_category_and_blocks_other_emails(self):
        with TemporaryDirectory() as directory:
            path = Path(directory) / 'truth.json'
            path.write_text(json.dumps({'email_001': {'category': 'BL_COMPARISON', 'defect_fields': ['secret']}, 'email_002': {'category': 'GENERAL'}}))
            with patch.dict(os.environ, {'EXTRACTION_DEBUG_GROUND_TRUTH': str(path)}):
                catalog = ExtractionCatalog()
                self.assertEqual(catalog.email_ids(), ['email_001'])
                inbox = Mock()
                with self.assertRaises(HTTPException):
                    catalog.email('email_002', inbox)
                inbox.emails.assert_not_called()
                api = TestClient(app)
                self.assertEqual(api.get('/api/v1/debug/extraction/emails').json(), {'email_ids': ['email_001']})
                with patch.object(InboxClient, 'emails') as emails:
                    self.assertEqual(api.post('/api/v1/debug/extraction', json={'email_id': 'email_002'}).status_code, 422)
                    emails.assert_not_called()
                self.assertEqual(api.post('/api/v1/debug/extraction', json={'email_id': '../x'}).status_code, 422)

    def test_unavailable_catalog(self):
        with patch.dict(os.environ, {'EXTRACTION_DEBUG_GROUND_TRUTH': '/missing/catalog.json'}):
            self.assertEqual(TestClient(app).get('/api/v1/debug/extraction/emails').status_code, 503)

    def test_api_fetches_only_selected_email_and_preserves_sides(self):
        email = Email(email_id='email_001', body='Compare docs', attachments=['attachments/email_001_BL.txt', 'attachments/email_001_SI.txt'])
        inbox = Mock()
        inbox.emails.return_value = [email]
        inbox.attachment.side_effect = [(BL, 'text/plain'), (SI, 'text/plain')]
        catalog = Mock()
        catalog.email.side_effect = lambda email_id, source: source.emails(email_id)[0]
        app.dependency_overrides[InboxClient] = lambda: inbox
        app.dependency_overrides[ExtractionCatalog] = lambda: catalog
        try:
            with patch.dict(os.environ, {'DEEPSEEK_API_KEY': ''}):
                response = TestClient(app).post('/api/v1/debug/extraction', json={'email_id': 'email_001'})
            self.assertEqual(response.status_code, 200)
            self.assertEqual(response.json()['bl']['fields']['shipper']['value'], 'Example Exporter')
            self.assertEqual(response.json()['si']['fields']['shipper']['value'], 'Example Exporter')
            self.assertEqual(response.json()['bl']['fields']['gross_weight_kg']['value'], '21577')
            self.assertEqual(response.json()['bl']['status'], 'extracted')
            self.assertEqual(response.json()['si']['status'], 'extracted')
            inbox.emails.assert_called_once_with('email_001')
        finally:
            app.dependency_overrides.clear()

    def test_pipeline_endpoint_uses_same_extractor_without_ground_truth(self):
        email = Email(email_id='email_001', body='Compare docs', attachments=['email_001_BL.txt', 'email_001_SI.txt'])
        inbox = Mock()
        inbox.emails.return_value = [email]
        inbox.attachment.side_effect = [(BL, 'text/plain'), (SI, 'text/plain')]
        app.dependency_overrides[InboxClient] = lambda: inbox
        try:
            with patch.dict(os.environ, {'EXTRACTION_DEBUG_GROUND_TRUTH': '/missing/catalog.json', 'DEEPSEEK_API_KEY': ''}):
                response = TestClient(app).post('/api/v1/extract', json={'email_id': 'email_001'})
            self.assertEqual(response.status_code, 200)
            self.assertEqual(response.json()['bl']['status'], 'extracted')
            self.assertEqual(response.json()['si']['status'], 'extracted')
            inbox.emails.assert_called_once_with('email_001')
        finally:
            app.dependency_overrides.clear()

    def test_partial_failure_and_missing_documents(self):
        email = Email(email_id='email_001', body='Compare', attachments=['email_001_BL.txt', 'email_001_SI.txt'])
        inbox = Mock()
        inbox.attachment.side_effect = [HTTPException(404, 'Missing BL'), (SI, 'text/plain')]
        result = extract_email(email, inbox)
        self.assertEqual(result.bl.status, 'error')
        self.assertEqual(result.si.status, 'extracted')
        email.attachments = ['first_BL.txt', 'second_BL.txt']
        inbox.reset_mock()
        result = extract_email(email, inbox)
        self.assertEqual(result.bl.status, 'needs_review')
        self.assertEqual(result.si.status, 'needs_review')
        inbox.attachment.assert_not_called()

    def test_wrong_document_and_missing_evidence_require_review(self):
        email = Email(email_id='email_001', body='Compare', attachments=['email_001_BL.txt', 'email_001_SI.txt'])
        inbox = Mock()
        inbox.attachment.side_effect = [(b'COMMERCIAL INVOICE\nGross Weight: 12 KG', 'text/plain'),
                                        (b'SHIPPING INSTRUCTION\nShipper: TBA\nGross Weight: 1200', 'text/plain')]
        result = extract_email(email, inbox)
        self.assertIsNone(result.bl.fields)
        self.assertEqual(result.bl.status, 'needs_review')
        self.assertIsNone(result.si.fields.shipper.value)
        self.assertEqual(result.si.fields.gross_weight_kg.value, '1200')
        self.assertIn('unit inferred', ' '.join(result.si.warnings))

    def test_shared_parser_handles_reference_aliases_and_conflicts(self):
        fields, warnings = extract_fields(BL.decode())
        self.assertFalse(warnings)
        self.assertEqual(fields.port_of_discharge.value, 'CALLAO')
        self.assertEqual(fields.container_count.value, '2')
        self.assertIn('Gross Wt (kgs): 21,577 KG', fields.gross_weight_kg.evidence)
        for name in type(fields).model_fields:
            evidence = getattr(fields, name).evidence
            self.assertIn(evidence, BL.decode())
        self.assertEqual(document_kind('BL INSTRUCTION\nShipper: Example'), 'SI')
        text = BL.decode() + 'Gross Weight (KG): 30,000 KG\n'
        fields, warnings = extract_fields(text)
        self.assertIsNone(fields.gross_weight_kg.value)
        self.assertIn('conflicting', ' '.join(warnings))

    def test_document_readers(self):
        self.assertEqual(read_document('SI.txt', b'Shipper: Example'), 'Shipper: Example')
        self.assertEqual(read_document('SI.txt', b'Shipper: Example\r\nConsignee: Other'), 'Shipper: Example\nConsignee: Other')
        self.assertEqual(read_document('SI.txt', 'Shipper: Caf\xe9'.encode('cp1252')), 'Shipper: Caf\xe9')
        document = Document()
        document.add_paragraph('SHIPPING INSTRUCTION')
        cells = document.add_table(rows=1, cols=2).rows[0].cells
        cells[0].text, cells[1].text = 'Gross Weight (KG)', '1200'
        data = BytesIO(); document.save(data)
        self.assertIn('Gross Weight (KG): 1200 KG', read_document('SI.docx', data.getvalue()))
        workbook = Workbook(); workbook.active.append(['Gross Weight (KG)', 1200])
        data = BytesIO(); workbook.save(data); workbook.close()
        self.assertIn('Gross Weight (KG): 1200 KG', read_document('BL.xlsx', data.getvalue()))
        writer = PdfWriter(); page = writer.add_blank_page(width=300, height=300)
        font = DictionaryObject({NameObject('/Type'): NameObject('/Font'), NameObject('/Subtype'): NameObject('/Type1'), NameObject('/BaseFont'): NameObject('/Helvetica')})
        page[NameObject('/Resources')] = DictionaryObject({NameObject('/Font'): DictionaryObject({NameObject('/F1'): font})})
        stream = DecodedStreamObject(); stream.set_data(b'BT /F1 12 Tf 10 100 Td (Shipper: Example) Tj ET')
        page[NameObject('/Contents')] = stream
        data = BytesIO(); writer.write(data)
        self.assertIn('Shipper: Example', read_document('BL.pdf', data.getvalue()))
        for path, content in [('image.png', b'123'), ('BL.pdf', b'corrupt'), ('SI.txt', b''), ('SI.docx', b'invalid')]:
            with self.subTest(path=path), self.assertRaises(DocumentReadError):
                read_document(path, content)

    def test_scanned_pdf_uses_ocr_and_text_pdf_skips_it(self):
        writer = PdfWriter()
        writer.add_blank_page(width=300, height=300)
        data = BytesIO(); writer.write(data)
        scanned = data.getvalue()
        with patch('app.services.extraction.pdf.ocr_page', return_value=BL.decode()) as ocr:
            self.assertIn('SHIPPER: Example Exporter', read_document('BL.pdf', scanned))
            ocr.assert_called_once_with(scanned, 1)

        writer = PdfWriter(); page = writer.add_blank_page(width=300, height=300)
        font = DictionaryObject({NameObject('/Type'): NameObject('/Font'), NameObject('/Subtype'): NameObject('/Type1'), NameObject('/BaseFont'): NameObject('/Helvetica')})
        page[NameObject('/Resources')] = DictionaryObject({NameObject('/Font'): DictionaryObject({NameObject('/F1'): font})})
        stream = DecodedStreamObject()
        stream.set_data(b'BT /F1 12 Tf 10 100 Td (Shipper: Example Exporter) Tj ET')
        page[NameObject('/Contents')] = stream
        data = BytesIO(); writer.write(data)
        with patch('app.services.extraction.pdf.ocr_page') as ocr:
            self.assertIn('Shipper: Example Exporter', read_document('BL.pdf', data.getvalue()))
            ocr.assert_not_called()

    def test_missing_ocr_binary_has_actionable_error(self):
        with patch('app.services.extraction.pdf.subprocess.run', side_effect=FileNotFoundError(2, 'missing', 'pdftoppm')):
            with self.assertRaisesRegex(DocumentReadError, 'Install Poppler and Tesseract'):
                ocr_page(b'pdf', 1)
