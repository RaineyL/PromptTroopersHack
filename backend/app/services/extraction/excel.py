"""Read XLSX rows as label/value lines without evaluating formulas."""
from io import BytesIO

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from .limits import DocumentReadError, MAX_TEXT, check_archive
from .tabular import format_cells


def read_excel(content: bytes) -> str:
    check_archive(content)
    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True, keep_links=False)
        try:
            lines = []
            length = 0
            for sheet in workbook:
                if (sheet.max_row or 0) > 10_000 or (sheet.max_column or 0) > 100:
                    raise DocumentReadError('Spreadsheet exceeds 10,000 rows or 100 columns.')
                for row in sheet.iter_rows(values_only=True):
                    line = format_cells([str(value) for value in row if value is not None])
                    if line:
                        lines.append(line)
                        length += len(line)
                        if len(lines) > 10_000 or length > MAX_TEXT:
                            raise DocumentReadError('Spreadsheet text exceeds the size limit.')
            return '\n'.join(lines)
        finally:
            workbook.close()
    except (InvalidFileException, KeyError, ValueError, OSError) as exc:
        if isinstance(exc, DocumentReadError):
            raise
        raise DocumentReadError('Spreadsheet could not be read. Check that the file is valid.') from exc
